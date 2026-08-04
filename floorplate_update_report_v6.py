
from __future__ import annotations

import argparse
import io
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

import fitz  # PyMuPDF
import pandas as pd
import pytesseract
from PIL import Image
from pptx import Presentation
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# Data model
# ============================================================

@dataclass
class UnitRecord:
    unit: str
    lot: Optional[int] = None
    internal: Optional[int] = None
    external: Optional[int] = None
    price: Optional[int] = None
    status: str = "Available"
    source_page: Optional[int] = None
    ppt_slide: Optional[int] = None


# ============================================================
# Normalisation helpers
# ============================================================

def normalize_unit(value: object) -> str:
    """
    Normalise:
      4.08   -> 408
      4 08   -> 408
      Unit408 -> 408
      1-G05  -> 1G05
    """
    text = str(value).strip().upper()
    text = re.sub(r"^UNIT\s*", "", text)
    return re.sub(r"[^A-Z0-9]", "", text)


def parse_int(value: object) -> Optional[int]:
    digits = re.sub(r"[^\d]", "", str(value))
    return int(digits) if digits else None


def parse_price(value: object) -> Optional[int]:
    return parse_int(value)


def is_available(value: object) -> bool:
    text = str(value).strip().lower()
    return text in {"available", "avail", "yes", "a"}


# ============================================================
# PPT parser
# ============================================================

UNIT_PRICE_PATTERN = re.compile(
    r"Unit\s*([A-Za-z0-9.\-]+)\s*-\s*\$\s*([\d,\s]+)",
    re.IGNORECASE,
)
AREA_PATTERN = re.compile(r"(\d+)\s*\+\s*(\d+)")


def parse_floorplate_pptx(path: str | Path) -> dict[str, UnitRecord]:
    prs = Presentation(str(path))
    records: dict[str, UnitRecord] = {}

    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            if "Unit" not in text:
                continue

            area_match = AREA_PATTERN.search(text)
            internal = int(area_match.group(1)) if area_match else None
            external = int(area_match.group(2)) if area_match else None

            for match in UNIT_PRICE_PATTERN.finditer(text):
                unit = normalize_unit(match.group(1))
                price = parse_price(match.group(2))
                if not unit or price is None:
                    continue

                records.setdefault(
                    unit,
                    UnitRecord(
                        unit=unit,
                        internal=internal,
                        external=external,
                        price=price,
                        ppt_slide=slide_number,
                    ),
                )

    return records


# ============================================================
# Excel price-list parser
# ============================================================

COLUMN_ALIASES = {
    "unit": ["unit", "unit number", "apartment", "apt"],
    "lot": ["lot", "lot number"],
    "internal": ["internal", "int", "internal size"],
    "external": ["external", "ext", "external size"],
    "price": ["price", "price aud", "price (aud)", "asking price"],
    "status": ["status", "availability"],
}


def _find_column(columns: Iterable[object], aliases: list[str]) -> Optional[str]:
    normalised = {str(c).strip().lower(): str(c) for c in columns}
    for alias in aliases:
        if alias in normalised:
            return normalised[alias]
    for normal, original in normalised.items():
        for alias in aliases:
            if alias in normal:
                return original
    return None


def parse_price_list_excel(path: str | Path, sheet_name: int | str = 0) -> dict[str, UnitRecord]:
    df = pd.read_excel(path, sheet_name=sheet_name)

    unit_col = _find_column(df.columns, COLUMN_ALIASES["unit"])
    lot_col = _find_column(df.columns, COLUMN_ALIASES["lot"])
    int_col = _find_column(df.columns, COLUMN_ALIASES["internal"])
    ext_col = _find_column(df.columns, COLUMN_ALIASES["external"])
    price_col = _find_column(df.columns, COLUMN_ALIASES["price"])
    status_col = _find_column(df.columns, COLUMN_ALIASES["status"])

    required = {
        "Unit": unit_col,
        "Internal": int_col,
        "External": ext_col,
        "Price": price_col,
    }
    missing = [name for name, col in required.items() if col is None]
    if missing:
        raise ValueError(f"Excel missing required columns: {', '.join(missing)}")

    records: dict[str, UnitRecord] = {}
    for _, row in df.iterrows():
        status = row[status_col] if status_col else "Available"
        if status_col and not is_available(status):
            continue

        unit = normalize_unit(row[unit_col])
        if not unit:
            continue

        records[unit] = UnitRecord(
            unit=unit,
            lot=parse_int(row[lot_col]) if lot_col else None,
            internal=parse_int(row[int_col]),
            external=parse_int(row[ext_col]),
            price=parse_price(row[price_col]),
            status=str(status),
        )

    return records


# ============================================================
# PDF price-list parser
# ============================================================

def _extract_pdf_text_rows(path: str | Path) -> list[tuple[int, str]]:
    """Use embedded text first."""
    doc = fitz.open(str(path))
    rows: list[tuple[int, str]] = []
    for page_number, page in enumerate(doc, start=1):
        text = page.get_text("text") or ""
        for line in text.splitlines():
            if line.strip():
                rows.append((page_number, line.strip()))
    return rows


def _ocr_pdf_rows(path: str | Path, dpi: int = 300) -> list[tuple[int, str]]:
    """
    OCR image-only PDFs.

    Streamlit Cloud requires:
      packages.txt: tesseract-ocr
      requirements.txt: pytesseract, pymupdf, pillow
    """
    doc = fitz.open(str(path))
    output: list[tuple[int, str]] = []

    for page_number, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png")))

        data = pytesseract.image_to_data(
            image,
            output_type=pytesseract.Output.DATAFRAME,
            config="--psm 6",
        )
        data = data.dropna(subset=["text"])
        data = data[data["conf"] >= 20]

        # Build complete visual rows by grouping nearby top positions.
        words = []
        for _, row in data.iterrows():
            text = str(row["text"]).strip()
            if not text:
                continue
            words.append(
                {
                    "text": text,
                    "left": int(row["left"]),
                    "top": int(row["top"]),
                    "height": int(row["height"]),
                }
            )

        words.sort(key=lambda x: (x["top"], x["left"]))
        grouped: list[list[dict]] = []
        tolerance = 18

        for word in words:
            centre_y = word["top"] + word["height"] / 2
            matched_group = None
            for group in grouped:
                group_y = sum(
                    item["top"] + item["height"] / 2 for item in group
                ) / len(group)
                if abs(centre_y - group_y) <= tolerance:
                    matched_group = group
                    break
            if matched_group is None:
                grouped.append([word])
            else:
                matched_group.append(word)

        for group in grouped:
            group.sort(key=lambda x: x["left"])
            line = " ".join(item["text"] for item in group)
            if line.strip():
                output.append((page_number, line.strip()))

    return output


def _parse_pdf_table_line(line: str, page_number: int) -> Optional[UnitRecord]:
    """
    Flexible row parser for rows resembling:
      4.08 66 4 2 2 0 1 77 m² 10 m² Yes N/W A$1,051,000 Available

    Supports:
      4.08 / 4 08 / 408
      A$ / $
      common OCR spacing variations
    """
    cleaned = (
        line.replace("m?", "m²")
        .replace("m2", "m²")
        .replace("A§", "A$")
        .replace("S/", "S/")
    )
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    # Must end in Available to avoid treating sold/reserved units as current.
    if not re.search(r"\bAvailable\b", cleaned, re.IGNORECASE):
        return None

    # Unit is expected at the start.
    unit_match = re.match(r"^\s*(\d{1,2}(?:[.\s]\d{1,2})|\d{3,4})\b", cleaned)
    if not unit_match:
        return None

    raw_unit = unit_match.group(1)
    remainder = cleaned[unit_match.end():].strip()

    price_match = re.search(r"A?\$\s*([\d,]+)", remainder, re.IGNORECASE)
    if not price_match:
        return None

    price = parse_price(price_match.group(1))

    # Extract all area-like values immediately followed by m².
    area_values = [
        parse_int(v)
        for v in re.findall(r"(\d{1,3})\s*m(?:²|\^?2|\?)", remainder, re.IGNORECASE)
    ]
    area_values = [v for v in area_values if v is not None]
    if len(area_values) < 2:
        return None

    internal, external = area_values[-2], area_values[-1]

    # Lot is normally the first integer after Unit.
    lot_match = re.match(r"^\s*([0-9\]\|Il]{2,3})\b", remainder)
    lot = None
    if lot_match:
        raw_lot = (
            lot_match.group(1)
            .replace("]", "1")
            .replace("|", "1")
            .replace("I", "1")
            .replace("l", "1")
        )
        lot = parse_int(raw_lot)

    return UnitRecord(
        unit=normalize_unit(raw_unit),
        lot=lot,
        internal=internal,
        external=external,
        price=price,
        status="Available",
        source_page=page_number,
    )


def parse_price_list_pdf(path: str | Path) -> dict[str, UnitRecord]:
    # Try embedded text.
    text_rows = _extract_pdf_text_rows(path)
    parsed = [
        record
        for page, line in text_rows
        if (record := _parse_pdf_table_line(line, page)) is not None
    ]

    # Image-only PDF fallback.
    if len(parsed) < 5:
        ocr_rows = _ocr_pdf_rows(path)
        parsed = [
            record
            for page, line in ocr_rows
            if (record := _parse_pdf_table_line(line, page)) is not None
        ]

    records: dict[str, UnitRecord] = {}
    for record in parsed:
        records.setdefault(record.unit, record)

    if not records:
        raise ValueError(
            "No Available units could be parsed from the PDF. "
            "Check that the PDF contains a table with Unit, INT, EXT, Price and Status."
        )

    return records


def parse_price_list(path: str | Path) -> dict[str, UnitRecord]:
    suffix = Path(path).suffix.lower()
    if suffix in {".xls", ".xlsx"}:
        return parse_price_list_excel(path)
    if suffix == ".pdf":
        return parse_price_list_pdf(path)
    raise ValueError("Price List must be XLS, XLSX or PDF.")


# ============================================================
# Comparison
# ============================================================

def compare_records(
    ppt_records: dict[str, UnitRecord],
    price_records: dict[str, UnitRecord],
) -> dict[str, list[dict]]:
    missing: list[dict] = []
    extra: list[dict] = []
    price_mismatch: list[dict] = []
    area_mismatch: list[dict] = []
    matched: list[dict] = []

    for unit, current in sorted(price_records.items()):
        old = ppt_records.get(unit)

        if old is None:
            missing.append(
                {
                    "Unit": unit,
                    "Lot": current.lot,
                    "Internal": current.internal,
                    "External": current.external,
                    "Price List Price": current.price,
                    "PDF Page": current.source_page,
                    "Action": "Add label to floorplate",
                }
            )
            continue

        price_same = old.price == current.price
        area_same = (
            old.internal == current.internal
            and old.external == current.external
        )

        if not price_same:
            price_mismatch.append(
                {
                    "Unit": unit,
                    "Lot": current.lot,
                    "PPT Price": old.price,
                    "Price List Price": current.price,
                    "Difference": (current.price or 0) - (old.price or 0),
                    "PPT Slide": old.ppt_slide,
                    "PDF Page": current.source_page,
                    "Action": "Update PPT price",
                }
            )

        if not area_same:
            area_mismatch.append(
                {
                    "Unit": unit,
                    "Lot": current.lot,
                    "PPT Internal": old.internal,
                    "PPT External": old.external,
                    "Price List Internal": current.internal,
                    "Price List External": current.external,
                    "PPT Slide": old.ppt_slide,
                    "PDF Page": current.source_page,
                    "Action": "Review area",
                }
            )

        if price_same and area_same:
            matched.append(
                {
                    "Unit": unit,
                    "Lot": current.lot,
                    "Internal": current.internal,
                    "External": current.external,
                    "Price": current.price,
                    "PPT Slide": old.ppt_slide,
                    "Result": "Match",
                }
            )

    for unit, old in sorted(ppt_records.items()):
        if unit not in price_records:
            extra.append(
                {
                    "Unit": unit,
                    "PPT Internal": old.internal,
                    "PPT External": old.external,
                    "PPT Price": old.price,
                    "PPT Slide": old.ppt_slide,
                    "Action": "Remove or verify availability",
                }
            )

    return {
        "missing": missing,
        "extra": extra,
        "price_mismatch": price_mismatch,
        "area_mismatch": area_mismatch,
        "matched": matched,
    }


# ============================================================
# Excel report builder
# ============================================================

NAVY = "17365D"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"

THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    title_cell = ws.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(bold=True, color=WHITE, size=16)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.font = Font(italic=True, color="5B6573", size=10)
    subtitle_cell.alignment = Alignment(horizontal="center")


def _write_table(
    ws,
    headers: list[str],
    rows: list[dict],
    fill_colour: Optional[str] = None,
    money_columns: Optional[list[str]] = None,
) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    if rows:
        for row_index, row in enumerate(rows, start=5):
            for col_index, header in enumerate(headers, start=1):
                cell = ws.cell(row_index, col_index, row.get(header))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = THIN_BORDER
                if fill_colour:
                    cell.fill = PatternFill("solid", fgColor=fill_colour)
                if money_columns and header in money_columns:
                    cell.number_format = '$#,##0;[Red]-$#,##0'
    else:
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(headers))
        cell = ws.cell(5, 1, "No issues found.")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A5"
    for index, header in enumerate(headers, start=1):
        width = max(12, min(32, len(header) + 4))
        ws.column_dimensions[get_column_letter(index)].width = width


def build_report(
    output_path: str | Path,
    project_name: str,
    ppt_name: str,
    price_list_name: str,
    ppt_records: dict[str, UnitRecord],
    price_records: dict[str, UnitRecord],
    comparison: dict[str, list[dict]],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    summary = {
        "Available units in Price List": len(price_records),
        "Units labelled in PPT": len(ppt_records),
        "Matched price and area": len(comparison["matched"]),
        "Price mismatches": len(comparison["price_mismatch"]),
        "Area mismatches": len(comparison["area_mismatch"]),
        "Missing from PPT": len(comparison["missing"]),
        "Extra in PPT": len(comparison["extra"]),
    }

    ws = wb.create_sheet("Summary")
    _style_title(ws, f"{project_name} - Floorplate Update Report", "Generated by V5", 6)

    headers = ["Check", "Count", "Action"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")

    rows = [
        ("Available units in Price List", summary["Available units in Price List"], "Source of truth"),
        ("Units labelled in PPT", summary["Units labelled in PPT"], "Existing labels"),
        ("Matched price and area", summary["Matched price and area"], "No action required"),
        ("Price mismatches", summary["Price mismatches"], "Update price"),
        ("Area mismatches", summary["Area mismatches"], "Review area"),
        ("Missing from PPT", summary["Missing from PPT"], "Add label"),
        ("Extra in PPT", summary["Extra in PPT"], "Remove or verify"),
    ]

    for row_index, values in enumerate(rows, 5):
        for col_index, value in enumerate(values, 1):
            cell = ws.cell(row_index, col_index, value)
            cell.border = THIN_BORDER
        colour = GREEN if row_index == 7 else (YELLOW if row_index >= 8 else None)
        if colour:
            for col_index in range(1, 4):
                ws.cell(row_index, col_index).fill = PatternFill("solid", fgColor=colour)

    ws.column_dimensions["A"].width = 31
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 26
    ws.freeze_panes = "A5"

    sheet_specs = [
        (
            "Missing in PPT",
            comparison["missing"],
            ["Unit", "Lot", "Internal", "External", "Price List Price", "PDF Page", "Action"],
            YELLOW,
            ["Price List Price"],
        ),
        (
            "Extra in PPT",
            comparison["extra"],
            ["Unit", "PPT Internal", "PPT External", "PPT Price", "PPT Slide", "Action"],
            RED,
            ["PPT Price"],
        ),
        (
            "Price Mismatch",
            comparison["price_mismatch"],
            [
                "Unit", "Lot", "PPT Price", "Price List Price", "Difference",
                "PPT Slide", "PDF Page", "Action",
            ],
            ORANGE,
            ["PPT Price", "Price List Price", "Difference"],
        ),
        (
            "Area Mismatch",
            comparison["area_mismatch"],
            [
                "Unit", "Lot", "PPT Internal", "PPT External",
                "Price List Internal", "Price List External",
                "PPT Slide", "PDF Page", "Action",
            ],
            ORANGE,
            [],
        ),
        (
            "Matched",
            comparison["matched"],
            ["Unit", "Lot", "Internal", "External", "Price", "PPT Slide", "Result"],
            GREEN,
            ["Price"],
        ),
    ]

    for sheet_name, rows_data, sheet_headers, fill_colour, money_cols in sheet_specs:
        ws = wb.create_sheet(sheet_name)
        _style_title(ws, sheet_name, project_name, len(sheet_headers))
        _write_table(ws, sheet_headers, rows_data, fill_colour, money_cols)

    # Full comparison.
    full_rows = []
    all_units = sorted(set(price_records) | set(ppt_records))
    for unit in all_units:
        current = price_records.get(unit)
        old = ppt_records.get(unit)

        if current and not old:
            result = "Missing in PPT"
        elif old and not current:
            result = "Extra in PPT"
        elif old.price != current.price:
            result = "Price mismatch"
        elif old.internal != current.internal or old.external != current.external:
            result = "Area mismatch"
        else:
            result = "Match"

        full_rows.append(
            {
                "Unit": unit,
                "Lot": current.lot if current else None,
                "PPT Internal": old.internal if old else None,
                "PPT External": old.external if old else None,
                "PPT Price": old.price if old else None,
                "Price List Internal": current.internal if current else None,
                "Price List External": current.external if current else None,
                "Price List Price": current.price if current else None,
                "PPT Slide": old.ppt_slide if old else None,
                "Result": result,
            }
        )

    ws = wb.create_sheet("Full Comparison")
    full_headers = [
        "Unit", "Lot", "PPT Internal", "PPT External", "PPT Price",
        "Price List Internal", "Price List External", "Price List Price",
        "PPT Slide", "Result",
    ]
    _style_title(ws, "Complete Unit Comparison", project_name, len(full_headers))
    _write_table(
        ws,
        full_headers,
        full_rows,
        None,
        ["PPT Price", "Price List Price"],
    )

    ws = wb.create_sheet("Notes")
    _style_title(ws, "Method and Source Notes", project_name, 6)
    notes = [
        ("PPT source", ppt_name),
        ("Price List source", price_list_name),
        ("Parser", "V5 complete-row PDF parser with OCR fallback"),
        ("Status rule", "Only Available rows are treated as current inventory"),
        ("Unit normalisation", "Dots, spaces and punctuation are removed before matching"),
        ("PPT modification", "This report function does not modify the PPT"),
    ]
    for row_index, (label, value) in enumerate(notes, start=4):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, value)
        ws.cell(row_index, 1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(row_index, 1).font = Font(bold=True)
        ws.cell(row_index, 1).border = THIN_BORDER
        ws.cell(row_index, 2).border = THIN_BORDER
        ws.cell(row_index, 2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 80

    wb.save(output_path)


# ============================================================
# Public API for Streamlit
# ============================================================

def generate_floorplate_update_report(
    pptx_path: str | Path,
    price_list_path: str | Path,
    output_path: str | Path,
    project_name: str = "Floorplate",
) -> dict[str, int]:
    ppt_records = parse_floorplate_pptx(pptx_path)
    price_records = parse_price_list(price_list_path)
    comparison = compare_records(ppt_records, price_records)

    build_report(
        output_path=output_path,
        project_name=project_name,
        ppt_name=Path(pptx_path).name,
        price_list_name=Path(price_list_path).name,
        ppt_records=ppt_records,
        price_records=price_records,
        comparison=comparison,
    )

    return {
        "available": len(price_records),
        "ppt_labelled": len(ppt_records),
        "matched": len(comparison["matched"]),
        "price_mismatch": len(comparison["price_mismatch"]),
        "area_mismatch": len(comparison["area_mismatch"]),
        "missing": len(comparison["missing"]),
        "extra": len(comparison["extra"]),
    }


def generate_report_from_uploaded_files(
    ppt_bytes: bytes,
    ppt_filename: str,
    price_bytes: bytes,
    price_filename: str,
    project_name: str = "Floorplate",
) -> tuple[bytes, dict[str, int]]:
    """
    Streamlit-friendly wrapper.

    Example:
        report_bytes, summary = generate_report_from_uploaded_files(
            ppt_file.getvalue(), ppt_file.name,
            price_file.getvalue(), price_file.name,
            project_name,
        )
    """
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        ppt_path = temp_dir_path / ppt_filename
        price_path = temp_dir_path / price_filename
        output_path = temp_dir_path / "Floorplate_Update_Report.xlsx"

        ppt_path.write_bytes(ppt_bytes)
        price_path.write_bytes(price_bytes)

        summary = generate_floorplate_update_report(
            pptx_path=ppt_path,
            price_list_path=price_path,
            output_path=output_path,
            project_name=project_name,
        )

        return output_path.read_bytes(), summary


# ============================================================
# CLI
# ============================================================

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Compare a floorplate PPTX with an XLS/XLSX/PDF Price List."
    )
    parser.add_argument("--floorplate", required=True, help="PPTX floorplate path")
    parser.add_argument("--price-list", required=True, help="XLS/XLSX/PDF price list path")
    parser.add_argument("--output", required=True, help="Output XLSX report path")
    parser.add_argument("--project-name", default="Floorplate")
    args = parser.parse_args()

    summary = generate_floorplate_update_report(
        pptx_path=args.floorplate,
        price_list_path=args.price_list,
        output_path=args.output,
        project_name=args.project_name,
    )
    print(summary)


if __name__ == "__main__":
    main()
