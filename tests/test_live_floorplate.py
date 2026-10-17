"""Synthetic records/contracts only; no company files, live requests or credentials."""
import copy
from dataclasses import replace
from datetime import datetime,timedelta,timezone
import io
import unittest
from unittest.mock import Mock,patch

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches,Pt
from live_floorplate.model import FloorplateError,snapshot,format_price,lot_key
from live_floorplate.api import sync,NoRedirect,origin
from live_floorplate.layout import parse_layout,matched
from live_floorplate.render import powerpoint,html_view,reference_style
from live_floorplate.sample import demonstration


def row(**changes):
    result = dict(project='Hurstville Rise',lotnumber='001',propertystatus='Available',
                  unit='G.01',bedrooms='2 Bed',bathrooms=2,internal=75,external=12,price='$1,000,000')
    result.update(changes)
    return result


def config():
    # A TEST contract, not a claimed getProject schema or authentication method.
    return dict(enabled=True,contract_verified=True,account_scope='agent_portal',method='GET',
                url='https://fixture.invalid/getProject',project='Hurstville Rise',pagination='next_link',
                rows_path='records',total_path='count',next_path='next',headers={},
                fields={k:k for k in row()})


class InventoryTests(unittest.TestCase):
    def test_exact_available_only(self):
        rows = [row(lotnumber=i,propertystatus=s) for i,s in enumerate(('Available','Held','Pending','Sold'),1)]
        data = snapshot(rows,4,'Hurstville Rise')
        self.assertEqual([a.lot for a in data.apartments],['1'])
        for status in ('available','Available pending','',None):
            with self.subTest(status=status),self.assertRaises(FloorplateError):
                snapshot([row(propertystatus=status)],1,'Hurstville Rise')

    def test_unavailable_rows_never_need_or_leak_price(self):
        data = snapshot([row(propertystatus='Sold',price='SENSITIVE_SENTINEL')],1,'Hurstville Rise')
        self.assertEqual(data.apartments,())

    def test_incomplete_duplicates_scope_and_bad_fields(self):
        for rows,total in (([row()],2),([row(),row(lotnumber=1)],2),([row(project='Other')],1),
                           ([row(price='POA')],1),([row(unit=None)],1),([row(bathrooms=None)],1),([row(internal='unknown')],1)):
            with self.subTest(rows=rows,total=total),self.assertRaises(FloorplateError):
                snapshot(rows,total,'Hurstville Rise')
        self.assertEqual(snapshot([],0,'Hurstville Rise').apartments,())

    def test_lot_matching_does_not_use_unit_or_order(self):
        sample,layout,pdf = demonstration()
        data = snapshot([row(lotnumber='DEMO-0-0',unit='CompletelyDifferentUnit')],1,'Hurstville Rise')
        self.assertEqual(matched(data,layout,0)[0][1].level,0)
        with self.assertRaises(FloorplateError):
            matched(snapshot([row(lotnumber='UNMAPPED')],1,'Hurstville Rise'),layout)
        self.assertEqual(lot_key('00014'),lot_key(14))
        for value in (True,14.0,None,'../14'):
            with self.assertRaises(FloorplateError): lot_key(value)

    def test_price_formatting_without_discount(self):
        for raw,want in (('A$1,051,000','$1,051,000'),('1000000.50','$1,000,000.50'),(1250000,'$1,250,000')):
            self.assertEqual(format_price(raw),want)
        for raw in (0,-1,'POA','NaN','1,00,000',True,'100.005'):
            with self.assertRaises(FloorplateError): format_price(raw)

    def test_stale_snapshot_blocks_both_views_and_export(self):
        data,layout,pdf = demonstration()
        stale = replace(data,synced_at=datetime.now(timezone.utc)-timedelta(minutes=6))
        for render in (lambda:html_view(stale,layout,pdf,0),lambda:powerpoint(stale,layout,pdf)):
            with self.assertRaises(FloorplateError): render()


class APITests(unittest.TestCase):
    def test_verified_request_context_needs_no_property_project_or_total(self):
        cfg=config()
        cfg.update(project_context_verified=True,completeness='expected_lot_set',expected_lots=[1],pagination='single',
                   query={'projectid':'SYNTHETIC_PROJECT'})
        cfg['fields'].pop('project')
        cfg.pop('total_path')
        item=row();item.pop('project')
        data=sync(cfg,Mock(return_value={'records':[item]}))
        self.assertEqual(data.source_rows,1)
        self.assertEqual(data.apartments[0].bathrooms,'2')
        cfg['project_context_verified']=False
        fetch=Mock()
        with self.assertRaises(FloorplateError):sync(cfg,fetch)
        fetch.assert_not_called()

    def test_countless_response_requires_exact_reviewed_lot_set(self):
        cfg=config();cfg.update(completeness='expected_lot_set',expected_lots=[1,2],pagination='single')
        with self.assertRaises(FloorplateError):sync(cfg,Mock(return_value={'records':[row()]}))
        cfg['expected_lots']=[]
        fetch=Mock()
        with self.assertRaises(FloorplateError):sync(cfg,fetch)
        fetch.assert_not_called()

    def test_signed_query_is_server_transport_only(self):
        cfg = config()
        cfg['pagination']='single'
        cfg['query']={'agencyid':'SYNTHETIC_AGENCY','sign':'SYNTHETIC_NOT_A_CREDENTIAL','projectid':'SYNTHETIC_PROJECT'}
        fetch=Mock(return_value=dict(records=[row()],count=1))
        data=sync(cfg,fetch)
        self.assertIn('sign=SYNTHETIC_NOT_A_CREDENTIAL',fetch.call_args.args[0])
        self.assertNotIn('SYNTHETIC_NOT_A_CREDENTIAL',repr(data))
        cfg['query']['projectid']='VERIFY_PROJECT_ID'
        fetch.reset_mock()
        with self.assertRaises(FloorplateError):sync(cfg,fetch)
        fetch.assert_not_called()

    def test_synthetic_pagination_complete(self):
        fetch = Mock(side_effect=[dict(records=[row()],count=2,next='/getProject?page=2'),
                                  dict(records=[row(lotnumber=2,propertystatus='Sold')],count=2,next=None)])
        data = sync(config(),fetch)
        self.assertEqual(data.source_rows,2)
        self.assertEqual(len(data.apartments),1)
        self.assertFalse(data.sample)
        self.assertEqual(fetch.call_args.args[0],'https://fixture.invalid/getProject?page=2')

    def test_missing_count_rows_auth_schema_and_changed_totals_fail(self):
        for payload in (dict(records=[row()],count=2,next=None),dict(records=[],count=None,next=None),
                        dict(records={},count=0,next=None),dict(login_required=True)):
            with self.subTest(payload=payload),self.assertRaises((FloorplateError,KeyError)):
                sync(config(),Mock(return_value=payload))
        with self.assertRaises(FloorplateError):
            sync(config(),Mock(side_effect=[dict(records=[row()],count=2,next='/next'),dict(records=[],count=3,next=None)]))
        with self.assertRaises(PermissionError): sync(config(),Mock(side_effect=PermissionError('REDACT_ME')))

    def test_redirects_and_cross_origin_pages_never_receive_headers(self):
        with self.assertRaises(FloorplateError):
            NoRedirect().redirect_request(None,None,302,'',{},'https://evil.invalid/')
        fetch = Mock(return_value=dict(records=[row()],count=2,next='https://evil.invalid/collect'))
        with self.assertRaises(FloorplateError): sync(config(),fetch)
        fetch.assert_called_once()
        for url in ('http://fixture.invalid/','https://user:pass@fixture.invalid/','https://fixture.invalid:0/'):
            with self.assertRaises(FloorplateError): origin(url)

    def test_unconfigured_api_never_calls_network(self):
        for change in (dict(enabled=False),dict(contract_verified=False),dict(account_scope='admin'),dict(method='POST'),
                       dict(pagination='unknown'),dict(headers={'Cookie':'SYNTHETIC_NOT_A_COOKIE'})):
            cfg = config();cfg.update(change)
            fetch = Mock()
            with self.assertRaises(FloorplateError): sync(cfg,fetch)
            fetch.assert_not_called()

    def test_repeated_pagination_and_duplicate_lots_fail(self):
        for fetch in (Mock(return_value=dict(records=[row()],count=2,next=config()['url'])),
                      Mock(side_effect=[dict(records=[row()],count=2,next='/next'),
                                        dict(records=[row()],count=2,next=None)])):
            with self.assertRaises(FloorplateError):sync(config(),fetch)


class RenderingTests(unittest.TestCase):
    def test_ppt_has_five_slides_and_only_available_labels(self):
        data,layout,pdf = demonstration()
        prs = Presentation(io.BytesIO(powerpoint(data,layout,pdf)))
        self.assertEqual(len(prs.slides),5)
        from live_floorplate.source_mapping import GROUPS
        for group,slide in enumerate(prs.slides):
            text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            for level in GROUPS[group][4]:
                self.assertIn(f'Unit DEMO-{level}00',text)
                for unavailable in (1,2,3): self.assertNotIn(f'Unit DEMO-{level}0{unavailable}',text)
            self.assertIn('2 Bed',text)
            self.assertIn('75m² + 12m²',text)
            self.assertIn('$1,000,000',text)
            self.assertEqual(sum(s.shape_type == 13 for s in slide.shapes),1)

    def test_html_and_ppt_share_labels_and_escape_unit(self):
        data,layout,pdf = demonstration()
        unsafe = replace(data,apartments=(replace(data.apartments[0],unit='<script>alert(1)</script>'),))
        html = html_view(unsafe,layout,pdf,0)
        self.assertNotIn('<script>',html)
        self.assertIn('&lt;script&gt;',html)
        self.assertNotIn('DEMO-001',html)
        self.assertIn('$1,000,000',html)

    def test_reference_style_reuses_v8_colored_label_detection(self):
        prs = Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(3),Inches(1))
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('F4C542')
        shape.text='Unit DEMO - $1'
        font=shape.text_frame.paragraphs[0].runs[0].font
        font.name='Arial';font.size=Pt(11);font.color.rgb=RGBColor(0,0,0)
        buffer=io.BytesIO();prs.save(buffer)
        style=reference_style(buffer.getvalue())
        self.assertEqual((style.fill,style.font,style.size),('F4C542','Arial',11))

    def test_layout_requires_all_levels_and_shared_pages(self):
        for pages in ({'0':0},{str(i):i for i in range(9)}):
            with self.assertRaises(FloorplateError):
                parse_layout(dict(version=1,project='Hurstville Rise',verified=True,pages=pages,positions=[]),9)


class ExistingToolsTests(unittest.TestCase):
    def test_v8_report_still_generates_from_synthetic_uploads(self):
        import pandas as pd
        from floorplate_update_report_v8 import generate_report_from_uploaded_files
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(3),Inches(1))
        box.fill.solid();box.fill.fore_color.rgb=RGBColor(255,255,0)
        box.text='75 + 12\nLot 1_Unit G01 - $1,000,000'
        ppt=io.BytesIO();prs.save(ppt)
        excel=io.BytesIO()
        pd.DataFrame([dict(Unit='G01',Lot=1,Internal=75,External=12,Price=1000000,Status='Available')]).to_excel(excel,index=False)
        report,summary=generate_report_from_uploaded_files(ppt.getvalue(),'synthetic.pptx',excel.getvalue(),'synthetic.xlsx','Synthetic')
        self.assertTrue(report.startswith(b'PK'))
        self.assertEqual(summary['available'],1)


class RefreshSafetyTests(unittest.TestCase):
    def test_api_failure_and_unmatched_refresh_clear_previous_prices(self):
        from live_floorplate.ui import synchronize_into,KEY
        data,layout,_=demonstration()
        state={KEY:data}
        with patch.dict('os.environ',{'LANDMARK_PREVIEW_ONLY':'0'}),patch('live_floorplate.ui.sync',side_effect=PermissionError('SENSITIVE_UPSTREAM_DETAIL')):
            with self.assertRaises(PermissionError): synchronize_into(state,config(),layout)
        self.assertNotIn(KEY,state)
        state[KEY]=data
        unmatched=snapshot([row(lotnumber='UNMAPPED')],1,'Hurstville Rise')
        with patch.dict('os.environ',{'LANDMARK_PREVIEW_ONLY':'0'}),patch('live_floorplate.ui.sync',return_value=unmatched):
            with self.assertRaises(FloorplateError): synchronize_into(state,config(),layout)
        self.assertNotIn(KEY,state)


class ResponseSuccessTests(unittest.TestCase):
    def test_configured_success_marker_rejects_error_envelope(self):
        cfg=config();cfg.update(success_path='success',success_value=True,pagination='single')
        for marker in (False,1,'true'):
            fetch=Mock(return_value={'success':marker,'records':[row()],'count':1})
            with self.assertRaises(FloorplateError):sync(cfg,fetch)
