"""115-Lot source evidence and five-group regressions, without company documents."""
import copy
import hashlib
import io
import json
from pathlib import Path
import unittest
import fitz
from PIL import Image
from pptx import Presentation
from live_floorplate.source_mapping import source_numbers,verify_source,GROUPS
from live_floorplate.layout import parse_layout
from live_floorplate.model import FloorplateError,snapshot
from live_floorplate.sample import source_preview
from live_floorplate.render import powerpoint,html_view,Style
from live_floorplate.grouping import information_boxes

MAPPING=json.loads((Path(__file__).resolve().parents[1]/'live_floorplate/hurstville_mapping.json').read_text())


def synthetic_source():
    # Reproduce numbered callout topology, never the original architectural artwork.
    mapping=copy.deepcopy(MAPPING);doc=fitz.open()
    image=io.BytesIO();Image.new('RGB',(100,100),'black').save(image,format='PNG')
    for page_index in range(6):
        page=doc.new_page(width=1190.55,height=841.89)
        if page_index:
            page.insert_image(fitz.Rect(400,300,1000,550),stream=image.getvalue())
        for row in mapping['positions']:
            if row['source_page']==page_index:
                bbox=row['source_bbox']
                page.insert_text((bbox[0],bbox[1]+7),row['lotnumber'],fontsize=7)
    pdf=doc.tobytes();doc.close()
    mapping['source_sha256']=hashlib.sha256(pdf).hexdigest()
    numbers=source_numbers(pdf)
    for row in mapping['positions']:row['source_bbox']=list(numbers[row['lotnumber']][1])
    boxes={}
    for row in mapping['positions']:boxes.setdefault(row['box_id'],[]).append(row)
    for box in boxes.values():
        x=sum((r['source_bbox'][0]+r['source_bbox'][2])/2 for r in box)/len(box)/1190.55
        y=sum((r['source_bbox'][1]+r['source_bbox'][3])/2 for r in box)/len(box)/841.89
        for row in box:row['anchor_x']=x;row['anchor_y']=y
    return mapping,pdf


class SourceMappingTests(unittest.TestCase):
    def test_checked_mapping_covers_all_lots_and_floors(self):
        self.assertEqual({r['lotnumber'] for r in MAPPING['positions']},{str(i) for i in range(1,116)})
        self.assertEqual(len(MAPPING['positions']),115)
        self.assertEqual(len({r['box_id'] for r in MAPPING['positions']}),57)
        for group,title,start,end,levels in GROUPS:
            rows=[r for r in MAPPING['positions'] if r['group']==group]
            self.assertEqual({int(r['lotnumber']) for r in rows},set(range(start,end+1)))
            for level in levels:self.assertEqual(sum(r['level']==level for r in rows),(end-start+1)//len(levels))
        self.assertFalse(MAPPING['review']['owner_publication_approval'])

    def test_synthetic_source_evidence_roundtrip(self):
        mapping,pdf=synthetic_source()
        self.assertTrue(verify_source(mapping,pdf))
        layout=parse_layout(mapping,6,pdf)
        self.assertEqual(len(layout.positions),115)

    def test_changed_pdf_missing_lot_bad_bbox_and_anchor_rejected(self):
        mapping,pdf=synthetic_source()
        mutations=[lambda m:m.update(source_sha256='wrong'),lambda m:m['positions'].pop(),
                   lambda m:m['positions'][0]['source_bbox'].__setitem__(0,0),
                   lambda m:m['positions'][0]['source_bbox'].__setitem__(0,float('nan')),
                   lambda m:m['positions'][0].update(anchor_x=float('nan')),
                   lambda m:m['positions'][0].update(anchor_x=.1),
                   lambda m:m['positions'][20].update(level=8)]
        for mutate in mutations:
            bad=copy.deepcopy(mapping);mutate(bad)
            with self.assertRaises(FloorplateError):verify_source(bad,pdf)
        with self.assertRaises(FloorplateError):parse_layout(mapping,6)

    def test_shared_boxes_use_identical_coordinates(self):
        boxes={}
        for row in MAPPING['positions']:boxes.setdefault(row['box_id'],[]).append(row)
        for box in boxes.values():
            self.assertEqual(len({tuple(r[k] for k in ('x','y','width','height','anchor_x','anchor_y')) for r in box}),1)

    def test_boxes_covering_artwork_or_other_boxes_rejected(self):
        mapping,pdf=synthetic_source()
        for xy in ((.5,.4),(mapping['positions'][1]['x'],mapping['positions'][1]['y'])):
            bad=copy.deepcopy(mapping);bad['positions'][0].update(x=xy[0],y=xy[1])
            with self.assertRaises(FloorplateError):verify_source(bad,pdf)

    def test_all_115_available_fit_exactly_five_slides(self):
        mapping,pdf=synthetic_source();layout=parse_layout(mapping,6,pdf)
        data=source_preview(layout)
        prs=Presentation(io.BytesIO(powerpoint(data,layout,pdf)))
        self.assertEqual(len(prs.slides),5)
        self.assertEqual(sum(len(information_boxes(data,layout,g)) for g in range(5)),57)
        for group,slide in enumerate(prs.slides):
            text='\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            for lot in range(GROUPS[group][2],GROUPS[group][3]+1):self.assertRegex(text,rf'Unit TEST{lot}(?:\s|$)')
            self.assertEqual(sum(s.shape_type==13 for s in slide.shapes),1)
            for shape in slide.shapes:
                if shape.shape_type==1:
                    self.assertEqual(shape.auto_shape_type,5) # Rounded rectangle
                    self.assertEqual(str(shape.fill.fore_color.rgb),'C2C7AD')

    def test_shared_box_omits_unavailable_prices_in_html_and_ppt(self):
        mapping,pdf=synthetic_source();layout=parse_layout(mapping,6,pdf)
        rows=[]
        for lot,status,price in ((19,'Available',900001),(33,'Held',900002),(47,'Available',900003)):
            rows.append(dict(project='Hurstville Rise',lotnumber=lot,propertystatus=status,unit=f'CHECK{lot}',
                             bedrooms='2 Bed',bathrooms=2,internal=75,external=12,price=price))
        data=snapshot(rows,3,'Hurstville Rise')
        boxes=information_boxes(data,layout,2)
        self.assertEqual(len(boxes),1)
        self.assertEqual(len(boxes[0][2]),2)
        html=html_view(data,layout,pdf,2)
        self.assertIn('CHECK19',html);self.assertIn('CHECK47',html)
        self.assertNotIn('CHECK33',html);self.assertNotIn('900,002',html)
        prs=Presentation(io.BytesIO(powerpoint(data,layout,pdf)))
        text='\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('2 Bath',text);self.assertNotIn('CHECK33',text);self.assertNotIn('900,002',text)
